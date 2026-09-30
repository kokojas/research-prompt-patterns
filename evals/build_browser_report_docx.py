#!/usr/bin/env python3
"""Build the Ukrainian local-analysis DOCX from frozen browser results."""

from __future__ import annotations

import argparse
import json
import random
import statistics
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = RGBColor(30, 50, 67)
BLUE = RGBColor(32, 91, 143)
MUTED = RGBColor(87, 101, 113)
PALE = "EAF1F6"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.bold = bold
            run.font.size = Pt(9)


def add_table(doc, heads, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(heads))
    table.style = "Table Grid"
    table.autofit = False
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    for i, head in enumerate(heads):
        set_cell_text(table.rows[0].cells[i], head, True)
        shade(table.rows[0].cells[i], PALE)
    for j, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], value)
            if j % 2:
                shade(cells[i], "F7F9FB")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text):
    return para(doc, text, "List Bullet")


def heading(doc, text, level=1):
    p = doc.add_heading(text, level)
    p.paragraph_format.keep_with_next = True
    return p


def figure(doc, path, caption, width=7.1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption, "Caption")
    cap.paragraph_format.keep_with_next = False
    return cap


def bootstrap_pair(records, arm, control, field, n=10000, seed=70419):
    groups = defaultdict(list)
    for row in records:
        groups[row["task_id"], row["arm"]].append(row)
    tasks = sorted({r["task_id"] for r in records})
    d = {t: statistics.mean(r[field] for r in groups[t, arm]) - statistics.mean(r[field] for r in groups[t, control]) for t in tasks}
    rng = random.Random(seed)
    draws = sorted(statistics.mean(d[rng.choice(tasks)] for _ in tasks) for _ in range(n))
    return statistics.mean(d.values()), (draws[int(.025*n)], draws[int(.975*n)-1])


def f2(x):
    return f"{x:+.3f}".replace(".", ",")


def ci_text(comparison):
    lo, hi = comparison["bootstrap_95"]
    return f"{f2(comparison['difference'])} [{f2(lo)}; {f2(hi)}]"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--analysis-dir", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    analysis = Path(args.analysis_dir)
    summary = json.loads((analysis / "local_score_summary.json").read_text(encoding="utf-8"))
    records = json.loads((analysis / "local_scored_records.json").read_text(encoding="utf-8"))
    audit = json.loads((analysis / "focused_audit.json").read_text(encoding="utf-8"))
    assert len(records) == 225 and audit["n_reviewed"] == 45
    arms = summary["arms"]
    comp = summary["comparisons"]
    req_v, req_v_ci = bootstrap_pair(records, "verify", "base", "required_coverage")
    req_c, req_c_ci = bootstrap_pair(records, "clarify", "two_turn_base", "required_coverage")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(.68)
    sec.top_margin = Inches(.66)
    sec.bottom_margin = Inches(.62)
    sec.header_distance = Inches(.32)
    sec.footer_distance = Inches(.32)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in [("Title", 22, NAVY), ("Heading 1", 14, BLUE), ("Heading 2", 11.5, NAVY)]:
        s = styles[name]
        s.font.name = "Arial"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        s.paragraph_format.space_after = Pt(5)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    if "Caption" not in styles:
        styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    styles["Caption"].font.name = "Arial"
    styles["Caption"].font.size = Pt(8.5)
    styles["Caption"].font.color.rgb = MUTED
    styles["Caption"].paragraph_format.space_after = Pt(8)

    header = sec.header.paragraphs[0]
    header.text = "ДОСЛІДЖЕННЯ ПРОМПТІВ  ·  ЛОКАЛЬНИЙ АНАЛІЗ"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(7.5)
        run.font.color.rgb = MUTED
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("30.09.2026  ·  стор. ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

    doc.add_paragraph("Оцінювання трьох дослідницьких промптів на 225 вебдіалогах", "Title")
    p = para(doc, "Локальний аналіз браузерного експерименту  ·  30 вересня 2026")
    p.runs[0].font.color.rgb = MUTED
    p.runs[0].font.size = Pt(10)
    heading(doc, "Підсумок", 1)
    para(doc, "П’ять умов пройшли по 15 різних вебзадачах із трьома незалежними повторами. Усі 225 фінальних діалогів завершені на GPT-5.6 Sol High. Дані оцінено з локальних файлів; жодні додаткові виходи Oracle не включено до аналізу.")
    bullet(doc, "Verification First: підтвердженого виграшу немає. Різниця у формальному індикаторі «критерії плюс пряме посилання» становить +0,050 [−0,054; +0,171], а покриття обов’язкових змістових маркерів — −0,017. Відповіді довші й повільніші.")
    bullet(doc, "Question Horizon: індикатор змістового розкриття прихованих аспектів зріс на +0,111 [ +0,037; +0,200], але обсяг відповіді збільшився з 697 до 2 197 слів, а час — з 67,0 до 148,2 с. Частина розширення корисна; цей індикатор не перевіряє істинність кожного твердження.")
    bullet(doc, "Clarify Then Investigate: перший раунд охопив значно більше релевантних тем для уточнення (+0,858 [ +0,777; +0,930]); середнє — 10,2 запитання. Покриття обов’язкових змістових маркерів у фіналі зросло на +0,049 [ +0,013; +0,090], нижче за наперед заданий поріг +0,10.")
    para(doc, "Висновок за зафіксованим протоколом: жодному промпту поки не можна присвоїти підтверджену перевагу у якості правильної, належно підтвердженої відповіді. Наявні числа — прозорий автоматизований скринінг; підтримку кожного суттєвого твердження джерелом і частоту матеріальних помилок у всіх 225 відповідях не сертифіковано.")

    heading(doc, "Дизайн і відтворюваність", 1)
    add_table(doc, ["Елемент", "Зафіксована умова"], [
        ("Вибірка", "15 задач: 4 фактологічні, 4 аналітичні, 4 практичні, 3 неоднозначні"),
        ("Повтори", "5 умов × 3 нові браузерні сесії на задачу = 225 діалогів"),
        ("Умови", "base; verify; horizon; two_turn_base; clarify"),
        ("Парні порівняння", "verify–base; horizon–base; clarify–two_turn_base"),
        ("Середовище", "Oracle browser engine; GPT-5.6 Sol; High; Web Search; concurrency 3"),
        ("Відсів", "Пілот 15 та попередні API-запуски не входять до фінальної вибірки"),
    ])
    para(doc, "У двоетапних умовах перше повідомлення містило короткий, навмисно неповний запит. Друге повідомлення надавало однакові повні деталі обом умовам. Воно не було відповіддю на кожне з 5–15 персоналізованих запитань, тому експеримент перевіряє поведінку уточнення й кінцевий результат за однакових фактів, але не повний ефект реальної взаємодії з користувачем.")
    para(doc, "Початково 25 запусків були технічно невдалими: 17 збоїв підтвердження Web Search і 8 тайм-аутів готовності. Під час повторів виникли ще 3 збої Web Search. Усього 28 невалідних спроб виключено; кожну замінено новою сесією з тими самими умовами. Для 225 фінальних відповідей у маніфестах підтверджено модель і High.")

    heading(doc, "Що саме виміряно", 1)
    para(doc, "Скринінг шукає фрази за рубрикою у фінальних відповідях. Дерево запитань Question Horizon вилучено перед оцінкою змісту.")
    para(doc, "У двоетапних діалогах запитання першого раунду відокремлено від URL із параметрами запиту, які спершу давали хибні підрахунки.")
    para(doc, "Для Verification First додатковий індикатор множить частку знайдених обов’язкових формулювань на наявність прямого URL із релевантного офіційного домену. Це інвентаризація посилань, не доказ, що сторінка підтверджує конкретне твердження. Через цю межу він не дорівнює передреєстрованій метриці grounded task completion.")
    para(doc, "Для невизначеності усереднено три повтори кожної пари «задача × умова», потім обчислено різницю на рівні 15 задач. 95% інтервали отримано 10 000 повторних вибірок задач. Вони не охоплюють похибку рубрики, вебджерел або автоматичного розпізнавання фраз.")

    heading(doc, "Основні числові результати", 1)
    add_table(doc, ["Умова", "Обов’язкові", "Приховані", "URL", "Слів разом", "Час, с"], [
        ("base", "0,968", "0,830", "41/45", "697", "67,0"),
        ("verify", "0,951", "0,822", "44/45", "879", "94,9"),
        ("horizon", "0,977", "0,941", "44/45", "2 197", "148,2"),
        ("two_turn_base", "0,940", "0,822", "43/45", "1 145", "99,4"),
        ("clarify", "0,988", "0,874", "44/45", "1 572", "116,8"),
    ])
    para(doc, "Частки «обов’язкові» й «приховані» — автоматичне покриття формулювань, а не людська оцінка правильності. «Слів разом» включає обидва раунди для двоетапних умов. Час включає виконання у браузері; токени внутрішнього міркування та фактичну вартість не спостерігалися.")
    add_table(doc, ["Парне порівняння", "Показник скринінгу", "Різниця та 95% інтервал", "Задачі + / = / −"], [
        ("verify / base", "формулювання + офіційний URL", ci_text(comp["verify"]), "5 / 6 / 4"),
        ("horizon / base", "приховані аспекти", ci_text(comp["horizon"]), "6 / 9 / 0"),
        ("clarify / two_turn_base", "теми першого раунду", ci_text(comp["clarify"]), "15 / 0 / 0"),
    ])
    figure(doc, analysis / "figures" / "primary-differences.png", "Рисунок 1. Різниці автоматизованих індикаторів. Рядки мають різні знаменники змісту і не є рейтингом трьох промптів.")
    para(doc, f"Перевірка чутливості для verify: без вимоги URL різниця обов’язкового покриття дорівнює {f2(req_v)} [{f2(req_v_ci[0])}; {f2(req_v_ci[1])}]. Це змінює знак ефекту і показує, наскільки висновок за «формулювання + URL» залежить від експорту посилань, особливо в P03/P04.")

    doc.add_page_break()
    heading(doc, "Де саме змінюється результат", 1)
    figure(doc, analysis / "figures" / "task-differences.png", "Рисунок 2. Зміни на рівні кожної з 15 задач. У стовпцях показано різні індикатори для зафіксованих пар умов.")
    para(doc, "Позитивний ефект Question Horizon за прихованими аспектами зосереджений у шести задачах, особливо F02 (версії WCAG 2.2), P03 (довірчі межі GitHub Actions), P04 (міграція пакування) та U02 (пріоритет після землетрусу). У дев’яти задачах скринінг не побачив різниці. Це сумісно з користю розширеної рамки там, де справді є приховані залежності; універсального приросту не видно.")
    para(doc, "Великий плюс verify в P04 частково пояснюється тим, що дві базові відповіді P04 містять лише внутрішні маркери цитат без прямих URL в експорті. Такий файл не дає перевірити джерело за прямим посиланням, але не доводить, що сам зміст відповіді неправильний.")

    doc.add_page_break()
    heading(doc, "Ціна додаткової структури", 1)
    figure(doc, analysis / "figures" / "efficiency.png", "Рисунок 3. Спостережуваний обсяг і час. Дві шкали подано окремо; це не оцінка фінансової вартості.")
    para(doc, "Порівняно з base, verify додає в середньому 182 слова і 27,9 с, Question Horizon — 1 501 слово і 81,1 с. Clarify проти свого двоетапного контролю додає 427 слів і 17,4 с. Отже, будь-яка змістова перевага повинна виправдовувати цю додаткову роботу та навантаження на читача.")
    figure(doc, analysis / "figures" / "length-vs-grounding.png", "Рисунок 4. Довжина відповіді та автоматично знайдені обов’язкові маркери. Скупчення біля 1,0 показує стелю індикатора; довші відповіді не можна автоматично вважати точнішими.")

    doc.add_page_break()
    heading(doc, "Перший раунд уточнень", 1)
    para(doc, "Clarify поставив у середньому 10,18 запитання; 44 із 45 діалогів містили 5–15 запитань. Виняток F03 поставив чотири й прямо пояснив, що п’яте суттєво не змінить дослідження. У two_turn_base середнє — 0,20: звичайна модель переважно відповідала на неповне питання відразу. У всіх 45 діалогах clarify перший раунд містив запитання, після чого надійшло друге повідомлення з умовою задачі.")
    para(doc, "Скринінг тем першого раунду: 0,952 для clarify проти 0,094 для двоетапного контролю. Перевага тут очікувана за самим форматом інструкції; вона не показує, що кожне запитання було необхідним. У простій фактологічній F01 про три стандарти NIST зразки ставили 11–12 запитань про інфраструктуру, терміни зберігання і регулювання, хоча заявлену в задачі фактологічну помилку можна було виправити без більшості цих відповідей. Це реальний ризик зайвого тертя. Натомість чотири запитання F03 про межі таймлайну та аудиторію були доречним винятком.")
    para(doc, f"За однакових повних фактів у другому повідомленні різниця фінального обов’язкового покриття становить лише {f2(req_c)} [{f2(req_c_ci[0])}; {f2(req_c_ci[1])}]. Це менше за передреєстрований поріг +0,10. У реальному застосуванні важливо дозволяти 5–15 запитань лише тоді, коли стільки невідомих справді впливають на рішення, і явно обробляти часткові відповіді користувача.")

    heading(doc, "Аудит оцінок і джерел", 1)
    para(doc, "Окремо перевірено 45 діалогів (20%): по дев’ять з кожної умови, з випадковими, категорійними, нетиповими та всіма дев’ятьма випадками без прямого URL у фінальному експорті. Точний список зафіксовано до виправлення правил скринінгу й опубліковано окремо, щоб пізніші зміни оцінок не підмінили вибірку. Переглянуто початок відповіді, свідчення за кожним критерієм, перелік доменів джерел і повний текст у спірних місцях. У дев’яти діалогах виправлено 11 позначок критеріїв. Зокрема, регулярний вираз іноді помилково зараховував тест колеса як перевірку в чистому середовищі або не розпізнавав явну відмову обирати місто лише за магнітудою.")
    para(doc, "Із 225 фінальних відповідей 216 містять прямий URL відповідного офіційного сімейства джерел. У дев’яти залишилися лише внутрішні маркери ChatGPT; дві з них мали URL в першому, але не фінальному раунді. Відсутній URL в експорті — це проблема відтворюваності цитування, а не автоматична фактологічна помилка. Кількість посилань або офіційний домен не гарантує підтримку конкретної тези.")
    para(doc, "Обмеження аудиту: він був сфокусованим, а не повним розбором усіх суттєвих тверджень 225 відповідей; етикетки умов у пакетах аудиту не були приховані. Це відхилення від зафіксованої вимоги сліпого оцінювання. Відповідно, частку суттєвих помилок та повний показник підтвердження джерелами не обчислено, а правило про не зростання помилок на 0,05 застосувати неможливо.")

    doc.add_page_break()
    heading(doc, "Практичне рішення", 1)
    add_table(doc, ["Промпт", "Що показав дослід", "Коли його доцільно використовувати"], [
        ("Verification First", "Немає переконливого приросту змісту; відповіді довші й повільніші.", "Коли важливий явний процес перевірки, але не очікувати доведеної переваги точності за цими даними."),
        ("Question Horizon", "Ширше розкриває приховані аспекти у частині задач, із великою ціною обсягу й часу.", "Для складного рішення з невідомими залежностями; вимагати короткого підсумку й перевірки джерел."),
        ("Clarify Then Investigate", "Систематично збирає контекст; фінальний змістовий приріст невеликий за наданого другого повідомлення.", "Коли бракує даних, що реально змінять рекомендацію; дозволити менше п’яти питань для простого фактчеку."),
    ])
    para(doc, "Найбільш захищений висновок: Question Horizon має ознаки користі для широти дослідження, Clarify надійно змінює поведінку першого раунду, а Verification First у цій вибірці не показав стійкого змістового покращення. Це не є рейтингом точності промптів. Для підтвердження потрібне сліпе семантичне оцінювання всіх контрольних критеріїв із посиланнями на конкретні твердження та незалежна перевірка суттєвих помилок.")

    heading(doc, "Межі перенесення висновку", 1)
    bullet(doc, "Одна модель, один рівень міркування, один браузерний інструмент і 15 задач; результати не узагальнюються автоматично на інші моделі або дати.")
    bullet(doc, "Скринінг має стелю біля 1,0, реагує на формулювання й не відрізняє правильну тезу від згадки хибної тези в запереченні чи цитаті.")
    bullet(doc, "Час браузерного запуску залежить також від пошуку та UI; він не є чистим часом моделі. Вартість API і приховані токени не спостерігались.")
    bullet(doc, "Дев’ять експортів із внутрішніми маркерами без URL ускладнюють оцінку джерел та спотворюють індикатор, який вимагає прямого посилання.")
    bullet(doc, "Другий раунд у clarify і контролі містив той самий повний текст задачі, але не відповіді на всі індивідуальні питання. Тому вартість запитань оцінити можна, а максимальний потенційний виграш від реального діалогу — ні.")

    heading(doc, "Дані та першоджерела", 1)
    para(doc, "Відтворювані файли: evals/BROWSER_WEB_PROTOCOL.md; evals/browser-web-tasks.json; evals/extract_browser_local.py; evals/score_browser_local.py; evals/record_browser_focused_audit.py. Локальні результати: local_records.json, local_score_summary.json, audit_sample.json і focused_audit.json у папці browser-web-local-analysis-20260930. Вихідний пакет browser-web-final-20260929 містить batch.json, design.json, tasks.json, маніфести та повні chat.md.")
    para(doc, "Контрольні офіційні джерела задач: NIST FIPS 203–205 — https://csrc.nist.gov/pubs/fips/203/final ; WCAG 2.2 — https://www.w3.org/TR/WCAG22/ ; NASA OSIRIS-REx — https://science.nasa.gov/mission/osiris-rex/ ; WMO State of the Global Climate 2024 — https://wmo.int/publication-series/state-of-global-climate/state-of-global-climate-2024 ; NOAA 2024 — https://www.ncei.noaa.gov/news/global-climate-202413 ; USGS magnitude/intensity — https://www.usgs.gov/faqs/what-difference-between-magnitude-and-intensity ; GitHub Actions security — https://docs.github.com/en/actions/security-for-github-actions/security-guides/security-hardening-for-github-actions ; PyPA setup.py — https://packaging.python.org/en/latest/discussions/setup-py-deprecated/ .")
    para(doc, "Джерела є опорними для рубрики; наведення URL у відповіді не дорівнює перевірці підтримки конкретної тези. Офіційні сторінки можуть оновлюватися після дати замороження набору задач 29.09.2026.")

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
